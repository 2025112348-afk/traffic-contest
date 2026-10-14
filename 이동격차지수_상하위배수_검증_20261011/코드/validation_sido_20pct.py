"""
이동격차지수 상·하위 20% 배수 검증 (시도별, 서울 제외)
========================================================
보고서 「이동격차지수 상·하위 격차 검증」 섹션의 모든 수치와 그림을 재현한다.

입력
  1) data/eupmyeon_master_3차.csv
     - Traffic_entropy_final.ipynb 에서 산출한 3차 이동격차지수 결과(eupmyeon_master)
     - 필요한 열: sidonm, sggnm, emdnm, 이동격차지수
  2) data/DRT_버스형_읍면동_상중.xlsx  (시트: 읍면동_DRT_상중)
     - 팀 자체 구축 DRT 운행 읍면동 현황 (지자체 홈페이지·고시·보도자료, 신뢰도 상·중)
     - 필요한 열: 시도, 시군구, 읍면동, 운영상태, 유형

방법
  - 서울 제외 3,128개 읍면동, 15개 시도 (전남·광주는 2026.7.1 통합 기준 한 단위)
  - 시도별로 지수 순 정렬 → 상위 20% / 하위 20% (같은 개수, 올림)
  - 배수 = 상위 평균 ÷ 하위 평균, 95% CI = 부트스트랩 5,000회
  - 시도 간 차이 = 순열검정 3,000회 (통계량: 시도별 log 배수의 표준편차)
  - 지수 타당성 = 상·하위 20%의 실제 DRT 운행 여부 Fisher 정확검정
  - 전국 10분위별 DRT 운행 비율 (정책제안 근거)
  ※ "상위 지수 > 하위 지수" 검정(Mann-Whitney 등)은 그룹을 지수로 나눴기 때문에
    항상 유의한 순환 검정이므로 사용하지 않는다.

출력 (output/)
  - 상하위배수_시도별20%_상위5.png   보고서 그림 (전국 + 격차 상위 4개 시도)
  - 지수10분위별_DRT비율.png          정책제안 그림
  - 이동격차지수_상하위배수_검증_시도별20%.xlsx

실행
  pip install pandas numpy scipy matplotlib openpyxl
  python validation_sido_20pct.py
  python validation_sido_20pct.py --master 경로.csv --drt 경로.xlsx --out output
"""
import argparse, math, re
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import fisher_exact
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

SEED = 20261011
N_BOOT = 5000
N_PERM = 3000
P = 0.2  # 상·하위 20%

# ------------------------------------------------------------------
# 0. 인자
# ------------------------------------------------------------------
ap = argparse.ArgumentParser()
ap.add_argument("--master", default="data/eupmyeon_master_3차.csv")
ap.add_argument("--drt", default="data/DRT_버스형_읍면동_상중.xlsx")
ap.add_argument("--drt-sheet", default="읍면동_DRT_상중")
ap.add_argument("--out", default="output")
args = ap.parse_args()
OUT = Path(args.out); OUT.mkdir(parents=True, exist_ok=True)
rng = np.random.default_rng(SEED)

# ------------------------------------------------------------------
# 1. 지역 키 정규화 (2026.7.1 행정구역 기준)
# ------------------------------------------------------------------
SIDO_ALIAS = {"광주광역시": "전남광주통합특별시", "전라남도": "전남광주통합특별시",
              "강원도": "강원특별자치도", "전라북도": "전북특별자치도"}
SHORT = {"강원특별자치도": "강원", "경기도": "경기", "경상남도": "경남", "경상북도": "경북",
         "대구광역시": "대구", "대전광역시": "대전", "부산광역시": "부산", "세종특별자치시": "세종",
         "울산광역시": "울산", "인천광역시": "인천", "전남광주통합특별시": "전남·광주",
         "전북특별자치도": "전북", "제주특별자치도": "제주", "충청남도": "충남", "충청북도": "충북"}

def norm_sido(s):
    s = str(s).strip(); return SIDO_ALIAS.get(s, s)

def norm_city(sido, sgg):
    """'수원시 장안구' → '수원시' (일반구를 시로 묶음), 세종은 단일."""
    if "세종" in sido:
        return "세종"
    return re.sub(r"^(\S+?시)\s*\S+구$", r"\1", str(sgg).strip())

def key(sido, sgg, emd):
    sido = norm_sido(sido)
    return f"{sido}|{norm_city(sido, sgg)}|{str(emd).strip()}"

# ------------------------------------------------------------------
# 2. 지수 마스터 로드 (서울 제외)
# ------------------------------------------------------------------
m = pd.read_csv(args.master, encoding="utf-8-sig")
m = m[~m.sidonm.str.contains("서울")].copy()
m["_k"] = [key(a, b, c) for a, b, c in zip(m.sidonm, m.sggnm, m.emdnm)]
assert m._k.is_unique, "읍면동 키 중복"
m["시도"] = m.sidonm.map(lambda s: SHORT[norm_sido(s)])
m["동읍면"] = np.where(m.emdnm.str.endswith(("동", "가")), "동", "읍면")
print(f"[지수] 서울 제외 {len(m):,}개 읍면동, {m.시도.nunique()}개 시도")

# ------------------------------------------------------------------
# 3. DRT 운행 읍면동 매칭
#    운영중·시범운영, 관광형 제외. 행정동 목록에 없는 법정동·생활권명은
#    후보 행정동 전부로 매핑한다(하위권 도시 동에 DRT가 늘어나는 방향 = 보수적).
# ------------------------------------------------------------------
d = pd.read_excel(args.drt, sheet_name=args.drt_sheet)
d = d[d.운영상태.isin(["운영중", "시범운영"]) & (d.유형 != "관광형") & (d.시도 != "서울특별시")].copy()

FIX = {("경기도", "고양시", "덕은동"): ["화전동"], ("경기도", "고양시", "향동동"): ["화전동"],
       ("경기도", "광주시", "고산동"): ["광남1동", "광남2동"], ("경기도", "광주시", "태전동"): ["광남1동", "광남2동"],
       ("경기도", "광주시", "장지동"): ["탄벌동"],
       ("경기도", "수원시", "권선동"): ["권선1동", "권선2동"], ("경기도", "수원시", "세류동"): ["세류1동", "세류2동", "세류3동"],
       ("경기도", "안양시", "안양9동"): ["병목안동"], ("경기도", "용인시", "유림동"): ["유림1동", "유림2동"],
       ("경기도", "의정부시", "민락동"): ["송산2동", "송산3동"],
       ("대구광역시", "북구", "산격동"): ["산격1동", "산격2동", "산격3동", "산격4동"],
       ("대구광역시", "수성구", "범물동"): ["범물1동", "범물2동"],
       ("대구광역시", "수성구", "대흥동"): ["고산1동", "고산2동", "고산3동"],
       ("대구광역시", "수성구", "연호동"): ["고산1동", "고산2동", "고산3동"]}

valid = set(m._k); drt_keys = set(); unmatched = []
for _, r in d.iterrows():
    k = key(r.시도, r.시군구, r.읍면동)
    if k in valid:
        drt_keys.add(k)
    elif (r.시도, r.시군구, r.읍면동) in FIX:
        drt_keys.update(key(r.시도, r.시군구, x) for x in FIX[(r.시도, r.시군구, r.읍면동)])
    else:
        unmatched.append((r.시도, r.시군구, r.읍면동))
assert drt_keys <= valid
m["DRT"] = m._k.isin(drt_keys).astype(int)
print(f"[DRT] 운행 읍면동 {m.DRT.sum()}곳 (매핑불가 {len(unmatched)}건)")

# ------------------------------------------------------------------
# 4. 상·하위 20% 배수, 부트스트랩 CI, DRT Fisher 검정
# ------------------------------------------------------------------
def ratio(v, k):
    s = np.sort(v); return s[-k:].mean() / s[:k].mean()

def boot_ci(v, k):
    n = len(v); idx = rng.integers(0, n, (N_BOOT, n))
    s = np.sort(v[idx], axis=1)
    r = s[:, -k:].mean(1) / s[:, :k].mean(1)
    return np.percentile(r, [2.5, 97.5])

def block(g, name):
    N = len(g); k = math.ceil(N * P)
    o = g.sort_values("이동격차지수", ascending=False)
    top, bot = o.head(k), o.tail(k)
    lo, hi = boot_ci(g.이동격차지수.to_numpy(), k)
    a, b = int(top.DRT.sum()), int(bot.DRT.sum())
    p = fisher_exact([[a, k - a], [b, k - b]])[1] if a + b > 0 else np.nan
    return dict(시도=name, 읍면동수=N, 상하위각=k,
                상위평균=top.이동격차지수.mean(), 하위평균=bot.이동격차지수.mean(),
                배수=top.이동격차지수.mean() / bot.이동격차지수.mean(), CI하=lo, CI상=hi,
                상위DRT=a, 하위DRT=b, 상위DRT율=a / k, 하위DRT율=b / k, Fisher_p=p)

rows = [block(m, "전국(서울 제외)")]
for s, g in m.groupby("시도"):
    rows.append(block(g, s))
R = pd.DataFrame(rows)
R = pd.concat([R.iloc[:1], R.iloc[1:].sort_values("배수", ascending=False)], ignore_index=True)

# ------------------------------------------------------------------
# 5. 시도 간 배수 차이 순열검정
# ------------------------------------------------------------------
def heterogeneity(df):
    groups = df.시도.to_numpy(); vals = df.이동격차지수.to_numpy(); labs = np.unique(groups)
    def stat(gr):
        lr = []
        for l in labs:
            v = vals[gr == l]; lr.append(math.log(ratio(v, math.ceil(len(v) * P))))
        return np.std(lr)
    obs = stat(groups)
    cnt = sum(stat(rng.permutation(groups)) >= obs for _ in range(N_PERM))
    return obs, (cnt + 1) / (N_PERM + 1)

het_stat, het_p = heterogeneity(m)

# ------------------------------------------------------------------
# 6. 전국 10분위별 DRT 운행 비율
# ------------------------------------------------------------------
m["분위"] = pd.qcut(m.이동격차지수.rank(method="first"), 10, labels=range(1, 11))
DEC = m.groupby("분위", observed=True).agg(
    읍면동=("DRT", "size"), DRT=("DRT", "sum"), 평균지수=("이동격차지수", "mean"),
    읍면비율=("동읍면", lambda s: (s == "읍면").mean()))
DEC["DRT율"] = DEC.DRT / DEC.읍면동
k_nat = math.ceil(len(m) * P)
top_nat = m.nlargest(k_nat, "이동격차지수")
no_drt = int((top_nat.DRT == 0).sum())

# ------------------------------------------------------------------
# 7. 콘솔 요약
# ------------------------------------------------------------------
pd.set_option("display.width", 200)
print("\n[시도별 상·하위 20% 배수]")
print(R[["시도", "읍면동수", "배수", "CI하", "CI상", "상위DRT", "하위DRT", "Fisher_p"]].round(3).to_string(index=False))
print(f"\n[시도 간 차이] 순열검정 통계량 {het_stat:.4f}, p = {het_p:.4f}")
print(f"[전국 상위 20%] DRT 없음 {no_drt}/{k_nat} = {no_drt / k_nat:.1%}")
print("[10분위 DRT율] " + ", ".join(f"{v:.1%}" for v in DEC.DRT율))

# ------------------------------------------------------------------
# 8. 그림
# ------------------------------------------------------------------
for f in ["Noto Sans CJK KR", "Malgun Gothic", "AppleGothic", "NanumGothic"]:
    if any(f == x.name for x in font_manager.fontManager.ttflist):
        plt.rcParams["font.family"] = f; break
plt.rcParams["axes.unicode_minus"] = False
INK, INK2, GRID, SURF, BLUE = "#0b0b0b", "#52514e", "#e7e6e2", "#fcfcfb", "#2a78d6"

# 8-1. 전국 + 격차 상위 4개 시도
top5 = R.head(5).iloc[::-1]
fig, ax = plt.subplots(figsize=(8.0, 3.6), facecolor=SURF); ax.set_facecolor(SURF)
y = np.arange(len(top5))
ax.errorbar(top5.배수, y, xerr=[top5.배수 - top5.CI하, top5.CI상 - top5.배수], fmt="o", color=BLUE,
            ecolor=BLUE, elinewidth=2.5, capsize=0, ms=8, mec=SURF, mew=1.5)
for xv, hi, yv in zip(top5.배수, top5.CI상, y):
    ax.text(hi + 0.1, yv, f"{xv:.1f}배", ha="left", va="center", fontsize=10, color=INK2)
labels = [f"{s}{' *' if s == '인천' else ''}  (n={n:,})" for s, n in zip(top5.시도, top5.읍면동수)]
ax.set_yticks(y); ax.set_yticklabels(labels, fontsize=10.5, color=INK)
ax.axhline(len(top5) - 1.5, color=GRID, lw=1)
ax.set_xlim(0, 7); ax.set_xlabel("상위 20% 평균 ÷ 하위 20% 평균 (배), 막대 = 부트스트랩 95% 신뢰구간", color=INK2, fontsize=9.5)
ax.tick_params(axis="y", length=0); ax.tick_params(axis="x", colors=INK2)
ax.grid(axis="x", color=GRID, lw=.8); ax.set_axisbelow(True)
for sp in ax.spines.values(): sp.set_visible(False)
fig.suptitle("시도별 이동격차지수 상위·하위 20% 배수 (격차 상위 4개 시도, 서울 제외)", x=0.03, ha="left",
             fontsize=12.5, color=INK, fontweight="bold")
fig.text(0.03, 0.885, "* 인천: 신설 구 결측이 0으로 계산돼 배수 과대 가능", ha="left", fontsize=8.5, color=INK2)
plt.tight_layout(rect=(0, 0, 1, 0.87))
plt.savefig(OUT / "상하위배수_시도별20%_상위5.png", dpi=200, facecolor=SURF); plt.close()

# 8-2. 10분위별 DRT 운행 비율
fig, ax = plt.subplots(figsize=(8.6, 3.8), facecolor=SURF); ax.set_facecolor(SURF)
x = np.arange(1, 11)
ax.bar(x, DEC.DRT율 * 100, width=0.62, color=BLUE)
for xi, v in zip(x, DEC.DRT율 * 100):
    ax.text(xi, v + 0.25, f"{v:.1f}%", ha="center", fontsize=8.5, color=INK)
ax.set_xticks(x); ax.set_xticklabels([str(i) for i in x], color=INK2)
ax.set_xlabel(f"이동격차지수 10분위 (1 = 가장 양호, 10 = 가장 취약), 서울 제외 {len(m):,}개 읍면동", color=INK2, fontsize=9.5)
ax.set_ylabel("DRT 운행 읍면동 비율 (%)", color=INK2, fontsize=9.5); ax.set_ylim(0, 16)
ax.tick_params(colors=INK2, length=0); ax.grid(axis="y", color=GRID, lw=.8); ax.set_axisbelow(True)
for sp in ax.spines.values(): sp.set_visible(False)
ax2 = ax.twiny(); ax2.set_xlim(ax.get_xlim()); ax2.set_xticks(x)
ax2.set_xticklabels([f"{v * 100:.0f}%" for v in DEC.읍면비율], fontsize=8, color=INK2); ax2.tick_params(length=0)
for sp in ax2.spines.values(): sp.set_visible(False)
ax2.set_xlabel("분위 내 읍·면 비율", fontsize=8.5, color=INK2)
fig.suptitle("지수 분위별 DRT 운행 비율: 6~7분위에서 가장 높고 최취약 분위에서 다시 낮아짐", x=0.03, ha="left",
             fontsize=12, color=INK, fontweight="bold")
plt.tight_layout(rect=(0, 0, 1, 0.93))
plt.savefig(OUT / "지수10분위별_DRT비율.png", dpi=200, facecolor=SURF); plt.close()

# ------------------------------------------------------------------
# 9. 엑셀
# ------------------------------------------------------------------
def table(df):
    return pd.DataFrame({
        "시도": df.시도.replace({"인천": "인천*", "세종": "세종*"}), "읍면동 수": df.읍면동수,
        "상·하위 각 읍면동 수": df.상하위각, "상위20% 평균지수": df.상위평균.round(3),
        "하위20% 평균지수": df.하위평균.round(3), "배수": df.배수.round(2),
        "95% CI 하한": df.CI하.round(2), "95% CI 상한": df.CI상.round(2),
        "DRT 상위(곳)": df.상위DRT, "DRT 하위(곳)": df.하위DRT,
        "DRT율 상위(%)": (df.상위DRT율 * 100).round(1), "DRT율 하위(%)": (df.하위DRT율 * 100).round(1),
        "Fisher p": df.Fisher_p.round(4).astype(object).where(df.Fisher_p.notna(), "-")})

dec_out = DEC.reset_index().assign(평균지수=lambda x: x.평균지수.round(3),
                                   읍면비율=lambda x: (x.읍면비율 * 100).round(1),
                                   DRT율=lambda x: (x.DRT율 * 100).round(1))
dec_out.columns = ["분위", "읍면동", "DRT 운행 읍면동", "평균지수", "읍면 비율(%)", "DRT율(%)"]
het_out = pd.DataFrame({"검정": [f"시도 간 배수 차이 (순열검정 {N_PERM:,}회, 통계량 = 시도별 log 배수 표준편차)"],
                        "관측 통계량": [round(het_stat, 4)], "p-value": [round(het_p, 4)]})

xlsx = OUT / "이동격차지수_상하위배수_검증_시도별20%.xlsx"
with pd.ExcelWriter(xlsx) as xw:
    table(R.head(5)).to_excel(xw, sheet_name="보고서_상위5", index=False)
    table(R).to_excel(xw, sheet_name="시도별_20%_전체", index=False)
    het_out.to_excel(xw, sheet_name="시도간_차이검정", index=False)
    dec_out.to_excel(xw, sheet_name="전국_10분위_DRT", index=False)

from openpyxl import load_workbook
from openpyxl.styles import Font
wb = load_workbook(xlsx)
for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    for c in ws[1]: c.font = Font(bold=True)
    for col in ws.columns:
        w = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(9, w * 1.6), 60)
for name in ("보고서_상위5", "시도별_20%_전체"):
    ws = wb[name]
    ws.cell(ws.max_row + 2, 1, "* 인천: 신설 4개 구(서해·검단·영종·제물포) 결측이 0으로 계산돼 배수 과대 가능 / "
                               "세종: 읍면동 24개로 표본 작음 / \"-\": 상·하위 모두 DRT 미확인")
wb.save(xlsx)
print(f"\n[완료] 결과 저장: {OUT.resolve()}")
